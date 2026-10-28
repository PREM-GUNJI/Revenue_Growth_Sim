from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend import db, export
from backend.api.main import app
from backend.assumptions.assumptions import SKU_IDS

SKU = "Aurora-can_330ml"
assert SKU in SKU_IDS
LEVER = {"price_index": 1.0, "promo_depth_pct": 0, "mechanic": "none", "promo_weeks_per_month": 0}


def scenario(name: str, price: float = 1.0, promo: int = 0) -> dict:
    lever = {**LEVER, "price_index": price, "promo_depth_pct": promo, "mechanic": "TPR" if promo else "none",
             "promo_weeks_per_month": 2 if promo else 0}
    return {"schema_version": 1, "name": name, "levers": {SKU: lever}, "cost_shock": {}}


@pytest.fixture()
def factory(monkeypatch: pytest.MonkeyPatch):
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    db.Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(db, "get_session_factory", lambda: session_factory)
    return session_factory


@pytest.fixture()
def client(factory) -> TestClient:  # noqa: ARG001
    return TestClient(app)


def make(client: TestClient, scenarios, name: str = "Q3 / pricing: review") -> str:
    return client.post("/workspaces", json={"name": name, "scenarios": scenarios}).json()["id"]


def deck_text(content: bytes) -> str:
    out = []
    for slide in Presentation(io.BytesIO(content)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                out.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(out)


def test_export_builds_a_deck_from_engine_results(client: TestClient) -> None:
    wid = make(client, [scenario("Baseline"), scenario("Promo 20%", promo=20), scenario("Price +4%", price=1.04)])
    response = client.post(f"/workspaces/{wid}/export")
    assert response.status_code == 200 and response.content[:2] == b"PK"
    assert response.headers["content-disposition"] == 'attachment; filename="Q3-pricing-review.pptx"'  # unsafe characters removed
    text = deck_text(response.content)
    assert "Q3 / pricing: review" in text and "Promo 20%" in text and "Price +4%" in text
    assert "Decision brief" in text and "Recommendation" in text and "Caveats" in text
    assert "synthetic" in text.lower()
    assert "7c3aed" not in text.lower()  # no purple-accent dependency in narrative copy


def test_every_percentage_in_the_deck_matches_the_engine() -> None:
    scenarios, results = export.evaluate([scenario("Baseline"), scenario("Promo 20%", promo=20)])
    rows = export.summary_rows(scenarios, results)
    expected = export._signed(rows[1]["gp_pct"])
    assert rows[1]["gp_pct"] is not None and rows[0]["gp_pct"] == pytest.approx(0)
    deck = export.build_deck("T", scenarios, results, "Tester")
    text = deck_text(deck)
    assert expected in text and rows[1]["scenario_id"][:16] in text
    story = export.executive_narrative(rows)
    assert story["recommendation"] is not None or story["hold_baseline"]


def test_refused_scenario_has_reasons_and_no_numbers() -> None:
    scenarios, results = export.evaluate([scenario("Baseline"), scenario("Wild", price=1.5)])
    rows = export.summary_rows(scenarios, results)
    assert rows[1]["status"] == "REFUSED" and rows[1]["reasons"]
    assert rows[1]["volume_pct"] is rows[1]["revenue_pct"] is rows[1]["gp_pct"] is None
    text = deck_text(export.build_deck("T", scenarios, results, "Tester"))
    assert "Wild" in text and "Caveats" in text and rows[1]["reasons"][0] in text
    assert "no figures" in text.lower() or "refused" in text.lower()
    # the refused scenario's name must not share a line with a percentage figure
    assert not any("%" in line and "Wild" in line for line in text.splitlines())


def test_executive_narrative_picks_guardrailed_winner_and_keeps_losers() -> None:
    scenarios, results = export.evaluate([
        scenario("Baseline"),
        scenario("Promo 20%", promo=20),
        scenario("Price +4%", price=1.04),
        scenario("Wild", price=1.5),
    ])
    story = export.executive_narrative(export.summary_rows(scenarios, results), guardrail_pct=5)
    assert story["refused"] and story["refused"][0]["name"] == "Wild"
    assert all(r["gp_pct"] is not None for r in story["options"])
    assert len(story["options"]) <= 3
    if story["recommendation"] is not None:
        assert story["recommendation"]["volume_pct"] >= -5
        assert story["recommendation"]["gp_pct"] > 0.05


def test_html_brief_is_sparse_and_matches_narrative() -> None:
    scenarios, results = export.evaluate([scenario("Baseline"), scenario("Promo 20%", promo=20), scenario("Wild", price=1.5)])
    page = export.build_brief("Case", scenarios, results, "Tester")
    assert "Decision brief" in page and "Recommended" in page and "Options at a glance" in page
    assert "Wild" in page and "refused" in page.lower()
    assert "<th>Result hash</th>" not in page  # hashes stay in the reproduce footer, not a table column


def test_refused_baseline_means_no_comparisons_at_all() -> None:
    scenarios, results = export.evaluate([scenario("Bad baseline", price=1.5), scenario("Promo", promo=20)])
    rows = export.summary_rows(scenarios, results)
    assert all(r["gp_pct"] is None and r["volume_pct"] is None for r in rows)


def test_ui_only_source_field_is_ignored() -> None:
    withsource = {**scenario("Baseline"), "source": {"label": "research", "candidate": {"price": 1}}}
    scenarios, _ = export.evaluate([withsource])
    assert scenarios[0].name == "Baseline"


def test_empty_and_invalid_workspaces_give_a_clear_422(client: TestClient) -> None:
    empty = client.post("/workspaces", json={"name": "Empty"}).json()["id"]
    assert client.post(f"/workspaces/{empty}/export").status_code == 422
    broken = make(client, [{"name": "No levers", "levers": {"nope": {"price_index": "x"}}}], "Broken")
    response = client.post(f"/workspaces/{broken}/export")
    assert response.status_code == 422 and "not valid" in response.json()["detail"]
    assert client.post("/workspaces/missing/export").status_code == 404


def test_export_is_audited(client: TestClient, factory) -> None:
    wid = make(client, [scenario("Baseline")])
    client.post(f"/workspaces/{wid}/export")
    with factory() as session:
        row = session.scalars(select(db.ActivityLog).where(db.ActivityLog.action == "workspace.export")).one()
    assert row.workspace_id == wid and row.detail == {"format": "pptx", "scenarios": 1}


@pytest.mark.real_auth
def test_export_requires_sign_in(factory) -> None:  # noqa: ARG001
    assert TestClient(app).post("/workspaces/x/export").status_code == 401
