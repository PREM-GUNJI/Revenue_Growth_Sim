"""PowerPoint export of a workspace. Every figure comes from `evaluate_batch`; nothing is estimated here.

A REFUSED scenario is listed with the engine's refusal reasons and no numbers at all. Percentages are
simple ratios of engine totals against the first scenario (the baseline), the same as the app's board.
"""

from __future__ import annotations

import io
from dataclasses import asdict
from datetime import UTC, datetime
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt

from backend.engine.batch import ENGINE_VERSION, _data_hash, evaluate_batch
from backend.engine.scenario import Scenario
from backend.model.spec import build_param_draws

INK = RGBColor(0x0A, 0x10, 0x34)
PURPLE = RGBColor(0x7C, 0x3A, 0xED)
GREY = RGBColor(0x52, 0x52, 0x52)
RED = RGBColor(0xDC, 0x26, 0x26)
K, SEED = 200, 42
SYNTHETIC = "All data in this simulator is synthetic. Figures are modelled, not observed."


class ExportError(ValueError):
    """The saved scenarios cannot be exported (empty, or not valid scenarios)."""


def _total(bands: dict[str, dict[str, float]]) -> float:
    return float(sum(band["value"] for band in bands.values()))


def _pct(value: float, base: float) -> float | None:
    return (value / base - 1) * 100 if base else None


def _signed(value: float | None) -> str:
    return "n/a" if value is None else f"{value:+.1f}%"


def evaluate(scenario_dicts: list[dict[str, Any]]) -> tuple[list[Scenario], list[dict[str, Any]]]:
    if not scenario_dicts:
        raise ExportError("This workspace has no scenarios to export.")
    try:
        scenarios = [Scenario.model_validate({k: v for k, v in item.items() if k != "source"}) for item in scenario_dicts]
    except Exception as exc:
        raise ExportError("Some saved scenarios are not valid and cannot be evaluated.") from exc
    results = [asdict(r) for r in evaluate_batch(scenarios, build_param_draws(k=K, seed=SEED))]
    return scenarios, results


def summary_rows(scenarios: list[Scenario], results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    base = results[0]
    base_ok = base["status"] != "REFUSED"
    rows = []
    for index, (scenario, result) in enumerate(zip(scenarios, results, strict=True)):
        row: dict[str, Any] = {"name": scenario.name, "status": result["status"], "scenario_id": result["scenario_id"],
                               "result_hash": result["result_hash"], "baseline": index == 0,
                               "reasons": [r["message"] for r in result["refusal_reasons"]]}
        if result["status"] == "REFUSED" or not base_ok:
            row.update(volume_pct=None, revenue_pct=None, gp_pct=None)  # refused: no numbers, ever
        else:
            row.update(volume_pct=_pct(_total(result["volume"]), _total(base["volume"])),
                       revenue_pct=_pct(_total(result["gsv"]), _total(base["gsv"])),
                       gp_pct=_pct(_total(result["gp"]), _total(base["gp"])))
        rows.append(row)
    return rows


def _text(slide, left, top, width, height, text, size=14, bold=False, color=INK):
    frame = slide.shapes.add_textbox(left, top, width, height).text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text, run.font.size, run.font.bold, run.font.color.rgb = text, Pt(size), bold, color
    return frame


def _add_note(frame, text, size=11, color=GREY):
    paragraph = frame.add_paragraph()
    run = paragraph.add_run()
    run.text, run.font.size, run.font.color.rgb = text, Pt(size), color


def build_deck(workspace_name: str, scenarios: list[Scenario], results: list[dict[str, Any]], author: str) -> bytes:
    rows = summary_rows(scenarios, results)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    width = deck.slide_width - Inches(1.2)

    cover = deck.slides.add_slide(blank)
    _text(cover, Inches(0.6), Inches(2.2), width, Inches(1.2), workspace_name, 38, True)
    _text(cover, Inches(0.6), Inches(3.4), width, Inches(0.6), "Revenue growth scenario report", 20, color=PURPLE)
    frame = _text(cover, Inches(0.6), Inches(4.6), width, Inches(1.4),
                  f"Prepared by {author} on {datetime.now(UTC):%d %b %Y}", 13, color=GREY)
    _add_note(frame, f"Engine {ENGINE_VERSION}, data {_data_hash()[:12]}, {K} sensitivity draws, seed {SEED}")
    _add_note(frame, SYNTHETIC, 12, RED)

    table_slide = deck.slides.add_slide(blank)
    _text(table_slide, Inches(0.6), Inches(0.4), width, Inches(0.7), "Scenarios against the baseline", 26, True)
    shape = table_slide.shapes.add_table(len(rows) + 1, 5, Inches(0.6), Inches(1.3), width, Emu(Inches(0.5) * (len(rows) + 1)))
    table = shape.table
    for col, head in enumerate(("Scenario", "Status", "Volume", "Revenue", "Gross profit")):
        table.cell(0, col).text = head
    for r, row in enumerate(rows, start=1):
        label = row["name"] + (" (baseline)" if row["baseline"] else "")
        values = (label, row["status"], _signed(row["volume_pct"]), _signed(row["revenue_pct"]), _signed(row["gp_pct"]))
        for col, value in enumerate(values):
            table.cell(r, col).text = "no figures: refused" if (row["status"] == "REFUSED" and col >= 2) else value
    for cell in (c for rw in table.rows for c in rw.cells):
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(13)

    chartable = [r for r in rows if r["gp_pct"] is not None and not r["baseline"]]
    if chartable:
        chart_slide = deck.slides.add_slide(blank)
        _text(chart_slide, Inches(0.6), Inches(0.4), width, Inches(0.7), "Gross profit change against baseline (%)", 26, True)
        data = CategoryChartData()
        data.categories = [r["name"] for r in chartable]
        data.add_series("Gross profit %", [round(r["gp_pct"], 2) for r in chartable])
        chart = chart_slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.3), width, Inches(5.4), data).chart
        chart.has_legend = False
        chart.plots[0].has_data_labels = True
        chart.plots[0].data_labels.number_format, chart.plots[0].data_labels.number_format_is_linked = '0.0"%"', False

    refused = [r for r in rows if r["status"] == "REFUSED"]
    if refused:
        slide = deck.slides.add_slide(blank)
        _text(slide, Inches(0.6), Inches(0.4), width, Inches(0.7), "Refused requests", 26, True)
        frame = _text(slide, Inches(0.6), Inches(1.3), width, Inches(5), "The engine will not extrapolate outside supported data, so these have no figures.", 14, color=GREY)
        for row in refused:
            _add_note(frame, row["name"], 16, INK)
            for reason in row["reasons"]:
                _add_note(frame, "    " + reason, 13)

    appendix = deck.slides.add_slide(blank)
    _text(appendix, Inches(0.6), Inches(0.4), width, Inches(0.7), "Traceability", 26, True)
    frame = _text(appendix, Inches(0.6), Inches(1.3), width, Inches(5.5), "Scenario ids and result hashes let every figure be reproduced by the engine.", 13, color=GREY)
    for row in rows:
        _add_note(frame, f"{row['name']}: {row['scenario_id'][:16]} / {row['result_hash'][:16]}", 12, INK)
    _add_note(frame, SYNTHETIC, 12, RED)

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()
